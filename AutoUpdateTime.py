import time
import xlwings as xw
import pandas as pd
from functools import reduce
from openpyxl import load_workbook
import traceback
import hashlib

# =========================
# CONFIGURATION
# =========================

MASTER_FILE = r"C:\Users\WarnerJo1\OneDrive - Carrier Corporation\Financials\Project Cost Analysis\Project Cost Analysis - Midwest.xlsx"
RAW_SHEET   = "Query"
DEST_FILE   = r"C:\Users\WarnerJo1\OneDrive - Carrier Corporation\Salesforce Imports\Regional Hours Upload to Smartsheet.xlsx"

DEST_SHEETS = {
    "Midwest": "Midwest",
    "North Central": "NorthCentral",
}

DISCIPLINES = ["Design Engineering", "Programming", "Graphics"]

# Refresh timeout safeguard (seconds)
REFRESH_TIMEOUT_SEC = 1800  # 30 minutes (configurable)

# =========================
# EXCEL REFRESH (DETERMINISTIC)
# =========================

def refresh_master_workbook(path: str, timeout_sec: int = REFRESH_TIMEOUT_SEC, visible: bool = True):
    """
    Open Excel, RefreshAll, wait for async queries + calculation to finish, then save/close.
    This eliminates the "needs manual open" symptom caused by async refresh timing.
    """
    print("▶ Refreshing master workbook …")

    app = xw.App(visible=visible)
    app.display_alerts = False
    app.screen_updating = False

    wb = None
    try:
        wb = app.books.open(path, update_links=False, read_only=False)

        # Trigger refresh
        wb.api.RefreshAll()
        print("• RefreshAll triggered.")

        # Wait for Power Query / async queries (if Excel exposes this)
        try:
            app.api.CalculateUntilAsyncQueriesDone()
            print("• Async queries finished.")
        except Exception:
            print("• Async wait API not available; using calc-state wait only.")

        # Wait until Excel finishes calculating
        # Excel states: 0=xlDone, 1=xlCalculating, 2=xlPending
        t0 = time.time()
        while app.api.CalculationState != 0:
            time.sleep(1)
            if time.time() - t0 > timeout_sec:
                raise TimeoutError(f"Excel refresh/calc did not finish within {timeout_sec} seconds.")

        # Strong finalization pass
        app.api.CalculateFullRebuild()

        wb.save()
        print("✅ Master refresh saved.")
    finally:
        if wb is not None:
            wb.close()
        app.quit()

# =========================
# HELPER FUNCTIONS
# =========================

def assign_region_from_prefix(project_name: str) -> str:
    """Classify region by reading the leading 2-letter prefix from the project string."""
    if project_name is None:
        return "Unknown"
    s = str(project_name).strip().upper()
    if len(s) < 2:
        return "Unknown"

    prefix = s[:2]  # first two characters

    north_central_prefixes = {"HA", "PI", "CV", "DT", "TL", "CB"}
    midwest_prefixes       = {"CH", "IA", "IN", "TC", "WI"}

    if prefix in north_central_prefixes:
        return "North Central"
    if prefix in midwest_prefixes:
        return "Midwest"
    return "Unknown"


def build_combined(df: pd.DataFrame) -> pd.DataFrame:
    df.columns = [c.strip() for c in df.columns]

    project_col      = "Details[project]"
    discipline_col   = "Details[*Task Description Grouped]"
    actual_hours_col = "[v_Details_All_Engineering_Hours]"
    est_hours_col    = "[v_Details_Budget_Hours]"

    required_cols = [project_col, discipline_col, actual_hours_col, est_hours_col]
    missing = [c for c in required_cols if c not in df.columns]
    if missing:
        raise ValueError(f"Expected columns not found: {missing}\nGot: {df.columns.tolist()}")

    def pivot_hours(source_col: str, label_prefix: str) -> pd.DataFrame:
        pivots = []
        for disc in DISCIPLINES:
            sub = df[df[discipline_col] == disc]
            piv = (
                sub.groupby(project_col, dropna=False)[source_col]
                   .sum()
                   .reset_index()
                   .rename(columns={source_col: f"{label_prefix}{disc} Hours"})
            )
            pivots.append(piv)
        return reduce(lambda L, R: pd.merge(L, R, on=project_col, how="outer"), pivots)

    actual_df = pivot_hours(actual_hours_col, "")
    est_df    = pivot_hours(est_hours_col, "Estimated ")

    combined = pd.merge(actual_df, est_df, on=project_col, how="outer")

    # Convert NaN → 0 for all hour columns (actual + estimated)
    hour_cols = [c for c in combined.columns if "Hours" in c]
    combined[hour_cols] = combined[hour_cols].fillna(0)

    # Optional: keep non-hour fields as None instead of NaN
    combined = combined.where(pd.notnull(combined), None)

    return combined


def update_table(ws, df: pd.DataFrame):
    # Read existing headers
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]
    has_headers = any(h for h in headers)

    if has_headers:
        # Normalize project header mapping if needed
        if "project" in headers and "Details[project]" in df.columns:
            df = df.rename(columns={"Details[project]": "project"})

        # Add any new df columns to header list
        for col in df.columns:
            if col not in headers:
                headers.append(col)

        # Rewrite header row
        for col_idx, col_name in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=col_name)

        # Ensure df has every header column
        for h in headers:
            if h not in df.columns:
                df[h] = None

        # Order df to match headers
        df = df[headers]

    else:
        # No headers yet: write them from df
        headers = list(df.columns)
        for col_idx, col_name in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=col_name)

    # Clear existing data rows (keep header row)
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    # Write data
    for row_idx, row in enumerate(df.itertuples(index=False), start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)


# ---- Meta sheet helpers (hidden, so we don't touch header row) ----

def get_meta_sheet(wb):
    name = "_meta"
    if name in wb.sheetnames:
        return wb[name]
    ws = wb.create_sheet(name)
    ws.sheet_state = "hidden"
    ws["A1"].value = "key"
    ws["B1"].value = "value"
    return ws

def get_meta_value(ws, key, default=""):
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == key:
            v = ws.cell(r, 2).value
            return default if v is None else v
    return default

def set_meta_value(ws, key, value):
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == key:
            ws.cell(r, 2).value = value
            return
    r = ws.max_row + 1
    ws.cell(r, 1).value = key
    ws.cell(r, 2).value = value


def df_fingerprint(df: pd.DataFrame) -> str:
    """
    Stable fingerprint of dataframe content to detect meaningful changes.
    Prevents false 'unchanged' when totals match but rows differ.
    """
    df2 = df.copy().fillna("")
    df2 = df2.astype(str)

    sort_cols = ["project"] if "project" in df2.columns else list(df2.columns)
    df2 = df2.sort_values(sort_cols).reset_index(drop=True)

    payload = df2.to_csv(index=False).encode("utf-8")
    return hashlib.md5(payload).hexdigest()

def read_sheet_as_df(ws) -> pd.DataFrame:
    """Read an openpyxl worksheet into a DataFrame using the first row as headers."""
    if ws.max_row < 1:
        return pd.DataFrame()
    headers = [cell.value for cell in ws[1]]
    if not any(h for h in headers):
        return pd.DataFrame()
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rows.append(row)
    return pd.DataFrame(rows, columns=headers)


def show_job_changes(region_name: str, old_df: pd.DataFrame, new_df: pd.DataFrame):
    """
    Compare old and new DataFrames by project and print which jobs changed,
    were added, or were removed, along with the specific column-level diffs.
    """
    hour_cols = [c for c in new_df.columns if "Hours" in str(c)]
    project_col = "project"

    if old_df.empty or project_col not in old_df.columns:
        if not new_df.empty:
            print(f"  [New data] {len(new_df)} project(s) written for the first time.")
        return

    old_df = old_df.copy()
    new_df = new_df.copy()

    # Normalize hour columns to float for comparison
    for col in hour_cols:
        if col in old_df.columns:
            old_df[col] = pd.to_numeric(old_df[col], errors="coerce").fillna(0)
        if col in new_df.columns:
            new_df[col] = pd.to_numeric(new_df[col], errors="coerce").fillna(0)

    old_projects = set(old_df[project_col].dropna().astype(str))
    new_projects = set(new_df[project_col].dropna().astype(str))

    added   = new_projects - old_projects
    removed = old_projects - new_projects
    common  = old_projects & new_projects

    changed_jobs = []
    for proj in sorted(common):
        old_row = old_df[old_df[project_col].astype(str) == proj].iloc[0]
        new_row = new_df[new_df[project_col].astype(str) == proj].iloc[0]
        diffs = []
        for col in hour_cols:
            old_val = float(old_row[col]) if col in old_df.columns else 0.0
            new_val = float(new_row[col]) if col in new_df.columns else 0.0
            if old_val != new_val:
                diffs.append(f"{col}: {old_val} → {new_val}")
        if diffs:
            changed_jobs.append((proj, diffs))

    print(f"\n  ── Job Change Report: {region_name} ──")
    if not added and not removed and not changed_jobs:
        print("  No job-level changes detected.")
        return

    if added:
        print(f"  NEW jobs ({len(added)}):")
        for p in sorted(added):
            print(f"    + {p}")

    if removed:
        print(f"  REMOVED jobs ({len(removed)}):")
        for p in sorted(removed):
            print(f"    - {p}")

    if changed_jobs:
        print(f"  CHANGED jobs ({len(changed_jobs)}):")
        for proj, diffs in changed_jobs:
            print(f"    ~ {proj}")
            for d in diffs:
                print(f"        {d}")
    print()


def finalize_workbook_for_shuttle(path: str, visible: bool = False):
    """
    Opens and saves the destination workbook in Excel to force:
      - UsedRange normalization
      - Excel "finalization" so cloud integrations (Data Shuttle) see the change immediately
    """
    print("▶ Finalizing destination workbook for Data Shuttle …")
    app = xw.App(visible=visible)
    app.display_alerts = False
    app.screen_updating = False

    wb = None
    try:
        wb = app.books.open(path, update_links=False, read_only=False)

        # Optional: harmless even if there are no formulas
        try:
            app.api.CalculateFullRebuild()
        except Exception:
            pass

        wb.save()
        print("✅ Destination workbook finalized (Excel open/save/close).")
    finally:
        if wb is not None:
            wb.close()
        app.quit()

# =========================
# MAIN
# =========================

def main():
    # 1) Refresh master deterministically
    # visible=True since you're watching; flip to False later for hands-off
    refresh_master_workbook(MASTER_FILE, visible=True)
    print("✅ Refresh complete.")

    # 2) Read refreshed raw data
    print("▶ Reading raw data …")
    df_raw = pd.read_excel(MASTER_FILE, sheet_name=RAW_SHEET, engine="openpyxl")
    df_raw.columns = df_raw.columns.str.strip()
    print(f"• Raw rows read from '{RAW_SHEET}': {len(df_raw)}")
    print(f"• Raw columns: {len(df_raw.columns)}")

    # Normalize discipline names (fix typo)
    discipline_col = "Details[*Task Description Grouped]"
    if discipline_col in df_raw.columns:
        df_raw[discipline_col] = df_raw[discipline_col].astype(str).str.strip().replace({
            "Progragmming": "Programming"
        })

    # Detect project column
    possible_project_cols = [
        "Details[project]", "Project", "Details[Project Name]",
        "Details[project name]", "Details[project_id]", "Details[Project]"
    ]
    project_col = next((c for c in possible_project_cols if c in df_raw.columns), None)
    if project_col is None:
        print("⚠️ Could not find a project column. Available columns are:")
        for c in df_raw.columns.tolist():
            print("  -", c)
        raise ValueError("Project column not found. Update 'possible_project_cols' to match your sheet.")
    else:
        print(f"• Using project column: {project_col}")
        print("• Sample project values:")
        print(df_raw[project_col].head(10).to_string(index=False))

    # Classify regions from 2-letter prefix
    df_raw["Details[region]"] = df_raw[project_col].apply(assign_region_from_prefix)

    # Diagnostics
    prefix_counts = df_raw[project_col].astype(str).str.slice(0, 2).str.upper().value_counts()
    print("• Prefix counts (first two letters):")
    print(prefix_counts.to_string())

    region_counts = df_raw["Details[region]"].value_counts(dropna=False)
    print("• Region classification counts:")
    print(region_counts.to_string())

    unknown_df = df_raw[df_raw["Details[region]"] == "Unknown"][[project_col]].copy()
    if not unknown_df.empty:
        unknown_path = "unknown_projects.csv"
        unknown_df.to_csv(unknown_path, index=False)
        print(f"⚠️ Saved {len(unknown_df)} 'Unknown' projects to {unknown_path} for inspection.")

    # 3) Open destination workbook
    dest_wb = load_workbook(DEST_FILE)
    meta_ws = get_meta_sheet(dest_wb)
    summary = []

    for region_name, sheet_name in DEST_SHEETS.items():
        print(f"▶ Processing {region_name} …")
        df_region = df_raw[df_raw["Details[region]"] == region_name].copy()
        print(f"• {region_name} raw rows: {len(df_region)}")

        ws = dest_wb[sheet_name] if sheet_name in dest_wb.sheetnames else dest_wb.create_sheet(sheet_name)
        prev_df = read_sheet_as_df(ws)

        if df_region.empty:
            print(f"• No rows for {region_name}. Clearing sheet '{sheet_name}'.")
            if ws.max_row > 1:
                ws.delete_rows(2, ws.max_row - 1)

            # Record empty state fingerprint so it doesn't repeatedly "update"
            set_meta_value(meta_ws, f"{sheet_name}_fingerprint", "EMPTY")
            summary.append(f"{region_name}: 0 projects")
            continue

        combined = build_combined(df_region)

        if "Details[project]" in combined.columns:
            combined.rename(columns={"Details[project]": "project"}, inplace=True)

        print(f"• {region_name} combined projects: {len(combined)}")

        fp = df_fingerprint(combined)
        key = f"{sheet_name}_fingerprint"
        prev_fp = get_meta_value(meta_ws, key, "")

        if prev_fp == fp:
            print(f"✓ {region_name}: unchanged (fingerprint match). Skipping update.")
            summary.append(f"{region_name}: unchanged")
        else:
            show_job_changes(region_name, prev_df, combined)
            update_table(ws, combined)
            set_meta_value(meta_ws, key, fp)
            print(f"⚠ Update detected for {region_name}: wrote {len(combined)} rows.")
            summary.append(f"{region_name}: updated ({len(combined)} rows)")

    # Stamp last run time (in hidden meta sheet)
    set_meta_value(meta_ws, "last_run", time.strftime("%Y-%m-%d %H:%M:%S"))

    dest_wb.save(DEST_FILE)

    # NEW: force Excel to "finalize" the saved workbook so Data Shuttle sees the update
    finalize_workbook_for_shuttle(DEST_FILE, visible=False)

    print("✅ All done.\nSummary:")
    for line in summary:
        print(" -", line)


if __name__ == "__main__":
    try:
        main()
    except Exception:
        print("⚠️ Error occurred:")
        traceback.print_exc()
    finally:
        input("\nPress Enter to exit…")